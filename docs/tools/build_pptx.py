"""Builds docs/exports/08_Presentation.pptx and docs/generated/slides_outline.md from the SAME slide list, using measured data.

Usage: backend/.venv/Scripts/python.exe docs/tools/build_pptx.py   (run before build_docs.py)"""
from __future__ import annotations

import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_docs as B  # noqa: E402  (loads the measured data D)

D, DOCS, META = B.D, B.DOCS, B.META
ACCENT = RGBColor(0x0B, 0x5F, 0xFF)
INK = RGBColor(0x11, 0x11, 0x11)
MUTED = RGBColor(0x55, 0x55, 0x55)

t, ev, bench, cov = D["tests"], D["bench"]["evaluation"], D["bench"], D["cov"]
ca = ev["conflict_detection"]
st = D["stats"]

# (title, bullets, image path relative to docs or None, speaker notes)
SLIDES = [
    ("IntelliMetrics AI", [
        "A governed semantic layer with TARA, an AI assistant for trusted business metrics",
        META['student'],
        f"Guide: {META['guide']} | {META['university']} | AY {META['year']}",
        "University project (academic capstone). Demonstration data is synthetic.",
    ], None, "Introduce yourself, the guide and the project name. Say at once that it is a university project with synthetic data."),
    ("The problem: one word, many numbers", [
        "Sales: Revenue = taxable value of goods sold (excluding GST)",
        "Finance: Revenue = sales less refunds",
        "Marketing: Revenue = invoice value including GST",
        "Nobody notices until two reports disagree; changes are not recorded",
        "Generic AI chatbots answer confidently even when wrong",
    ], None, "Use the Indian GST example because the audience recognises it. Stress that nobody is lying; the same word means different calculations."),
    ("Objectives", [
        "One governed live definition per business metric",
        "Detect conflicting definitions automatically and measure the detector",
        "Immutable version history with safe rollback",
        "Let non-technical users ask questions without risking wrong or unsafe queries",
        "Localise to India: INR (lakh/crore), April-March FY, GST, festivals, IST",
    ], None, "Map each objective to a result later in the talk."),
    ("The solution at a glance", [
        "Datasets in, governed metrics and dashboards out",
        "Workflow: draft, review, approve, publish (exactly one live per metric)",
        "Conflict centre: assign, review, approve, resolve",
        "Formula Lab: safe, parameterised execution",
        "Ask TARA: plain-English questions answered from approved metrics",
        f"{st['endpoints']} API endpoints | {st['pages']} pages | 4 roles | 2 isolated workspaces",
    ], "diagrams/14_use_case_roles.png", "Give the one-minute product tour using the use-case diagram: four roles, each with a distinct job."),
    ("Architecture", [
        "React + TypeScript single-page app",
        "FastAPI back end: RBAC on every route, workspace filter on every query",
        "PostgreSQL / SQLite for metadata, DuckDB + Parquet for analytics",
        "Celery + Redis for background ingestion (thread fallback)",
    ], "diagrams/01_high_level_architecture.png", "Explain that analytics never touch the transactional database; DuckDB reads Parquet files produced by ingestion."),
    ("Governance: versions that cannot be changed", [
        "Partial unique index: one live definition per metric key",
        "Version rows immutable: ORM guard + database trigger",
        "Rollback creates a NEW version; history is never rewritten",
        "Dashboards follow latest published or pinned versions",
    ], "diagrams/08_metric_state_machine.png", "Say that the guarantee is in the database, so a bug in the application cannot break it. Verified on real PostgreSQL 18."),
    ("A formula engine that cannot be injected", [
        "Lexer, recursive-descent parser, typed syntax tree",
        "Whitelisted functions; columns checked against the schema",
        "Every value is a bound parameter; division is zero-safe",
        "Limits: 2,000 characters, depth 40, 10,000 rows, 20 s timeout",
        "25 hostile inputs tested: all rejected before execution",
    ], "diagrams/06_formula_compile_pipeline.png", "No user text ever becomes SQL text. That is why injection is structurally impossible rather than filtered."),
    ("Automatic conflict detection", [
        "Compares formula, aggregation, filters, unit, time window, dataset",
        "Severity score, side-by-side diff, recommended resolution",
        f"{ca['cases']} labelled cases: precision {ca['automated']['precision']:.2f}, recall {ca['automated']['recall']:.2f}",
        f"Plain text comparison baseline: precision {ca['naive_manual']['precision']:.2f}, recall {ca['naive_manual']['recall']:.2f}",
    ], "diagrams/07_conflict_detection_flow.png", "Be candid: the cases were written by me, so this shows the rules behave as designed and is not an independent benchmark."),
    ("Ask TARA: the AI never writes numbers", [
        "Question becomes a plan; the plan is validated against approved metrics",
        "Same safe compiler runs it; the answer text is composed from rows",
        "Compare ranges, second dimension, share of total, default fallbacks",
        "Refuses when the data cannot answer",
    ], "diagrams/04_ask_tara_request_flow.png", "Walk through the sequence diagram. Mention the offline deterministic mode and that no live AI key was used."),
    ("Built for India", [
        "INR with lakh/crore grouping; compact forms ₹45 L, ₹1.2 Cr",
        "Financial year 1 April to 31 March; fiscal quarters",
        "GST: taxable value vs invoice value; slabs 0/5/12/18/28 %",
        "Festival windows: Diwali, Dussehra, Pongal, Eid, Holi, End-of-Season sale",
        "DD-MMM-YYYY dates in IST",
    ], "diagrams/12_fiscal_calendar_and_planner.png", "Explain that 'last quarter' resolves to the last completed fiscal quarter, and 'this year' to the financial year to date."),
    ("Demonstration", [
        "Reviewer: detect and resolve the Revenue conflict",
        "Analyst: Formula Lab, including a rejected injection",
        "Business user: Diwali comparison and UPI vs COD split",
        "Admin: audit log; other workspace: no cross-company data",
    ], "screenshots/light/41_ask_tara_empty.png", "Run the live demo here (8 minutes). The screenshot is a backup if the demo fails."),
    ("Testing", [
        f"{t['total']} automated tests, {t['failed']} failed",
        f"pytest {t['pytest']['total']} (+ {t['pytest_pg']['total']} on PostgreSQL 18), Vitest {t['vitest']['total']}, Playwright {t['playwright']['total']}",
        f"Back-end statement coverage {cov['backend'].get('pct', 'n/a')} %",
        "Security tests: injection, token forgery, CSRF, isolation, escalation",
        "Every page screenshotted in light and dark mode",
    ], "charts/tests_by_suite.png", "Each row of the test-case table in the report is linked to a measured result, so the table cannot claim a pass that did not happen."),
    ("Measured performance", [
        f"Query engine median {bench['query_latency']['engine_ms']['p50']} ms, p95 {bench['query_latency']['engine_ms']['p95']} ms on {bench['dataset']['rows']:,} rows",
        f"Ingestion about {ev['dataset_processing']['rows_per_second']:,} rows per second",
        f"Dashboard load {ev['dashboard_load']['avg_ms']} ms on average",
        "Time-saving estimates rely on stated assumptions, not measurements",
    ], "charts/query_latency.png", "Single-user numbers from one Windows machine. Say clearly that there was no load test."),
    ("Honest limitations", [
        "Docker Compose written but never run (Docker not installed)",
        "No live AI provider tested; offline planner only",
        "Synthetic data; no user study; no external security audit",
        "Chromium only; no load test; per-process rate limiter",
        "Verified instead: PostgreSQL 18.2, Redis + Celery, MySQL 8.0",
    ], None, "Say these before the examiner asks. It builds trust."),
    ("Future scope and conclusion", [
        "Evaluate with a real LLM on a held-out question set",
        "Redis-backed rate limiting; e-mail delivery; single sign-on",
        "Column-level security; connectors to real BI sources; user study",
        "One definition, one history, answers you can trace",
    ], None, "Close with the contribution: an integrated, tested, honest implementation."),
    ("Thank you. Questions?", [
        META['student'],
        f"Guide: {META['guide']}",
        "Documents: Project Report, Concept Note, Architecture, Testing Report, User Guide, Deployment Guide, Viva Preparation",
    ], None, "Invite questions. The Viva Preparation document lists forty likely ones."),
]


def add_text(slide, left, top, width, height, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, "Calibri"
    return box


def fit_picture(slide, path: Path, left, top, width, height):
    im = Image.open(path)
    w, h = im.size
    scale = min(width / w, height / h)
    pw, ph = int(w * scale), int(h * scale)
    slide.shapes.add_picture(str(path), left + (width - pw) // 2, top + (height - ph) // 2, pw, ph)


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    outline = ["# Slide-by-slide outline\n", "The deck `exports/08_Presentation.pptx` contains these slides, with the speaker notes below in the PowerPoint notes pane.\n"]
    for i, (title, bullets, img, notes) in enumerate(SLIDES, 1):
        s = prs.slides.add_slide(blank)
        W, H = prs.slide_width, prs.slide_height
        bar = s.shapes.add_shape(1, 0, 0, Inches(0.25), H)
        bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
        first = i == 1
        add_text(s, Inches(0.7), Inches(0.45 if not first else 2.0), Inches(12), Inches(1.1), title, 44 if first else 32, True)
        has_img = img and (DOCS / img).exists()
        body_w = Inches(6.3) if has_img else Inches(11.8)
        top = Inches(1.6 if not first else 3.3)
        box = s.shapes.add_textbox(Inches(0.7), top, body_w, Inches(5.2))
        tf = box.text_frame
        tf.word_wrap = True
        for j, b in enumerate(bullets):
            p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = b if first else "•  " + b
            r.font.size, r.font.name, r.font.color.rgb = Pt(20 if first else 19), "Calibri", (MUTED if first else INK)
            p.space_after = Pt(10)
        if has_img:
            fit_picture(s, DOCS / img, Inches(7.2), Inches(1.5), Inches(5.7), Inches(5.3))
        add_text(s, Inches(0.7), Inches(7.0), Inches(12), Inches(0.35), f"{META['student']} | {META['university']} | University project, synthetic demonstration data | {i}/{len(SLIDES)}", 10, False, MUTED)
        s.notes_slide.notes_text_frame.text = notes
        outline.append(f"\n## Slide {i}: {title}\n")
        outline += [f"* {b}" for b in bullets]
        if has_img:
            outline.append(f"\n*Visual:* `{img}`")
        outline.append(f"\n*Speaker notes:* {notes}")
    prs.core_properties.author = META["student"]
    prs.core_properties.title = "IntelliMetrics AI: Presentation"
    out = B.EXPORTS / "08_Presentation.pptx"
    prs.save(str(out))
    (B.GEN / "slides_outline.md").write_text("\n".join(outline) + "\n", encoding="utf-8")
    print("wrote", out, f"{out.stat().st_size // 1024} KB;", len(SLIDES), "slides")


if __name__ == "__main__":
    build()
