"""Tabular exports to CSV / XLSX / PDF. Each report is (title, columns, rows)."""
from __future__ import annotations

import csv
import io
from datetime import datetime

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..models import Conflict, Dashboard, EvaluationRun, Metric, User
from .dashboards import dashboard_data
from .formatting import fmt_datetime_ist, format_value

Report = tuple[str, list[str], list[list]]


def metrics_report(db: Session, user: User) -> Report:
    rows = db.scalars(select(Metric).where(Metric.workspace_id == user.workspace_id).order_by(Metric.name)).all()
    cols = ["Name", "Key", "Status", "Department", "Category", "Unit", "Formula", "Filters", "Dataset", "Owner", "Version", "Published version", "Updated"]
    data = [[m.name, m.key, m.status, m.department, m.category, m.unit, m.formula, m.filters, m.dataset.name if m.dataset else "", m.owner.name if m.owner else "", m.version, m.published_version, fmt_datetime_ist(m.updated_at)] for m in rows]
    return "Metrics catalog", cols, data


def conflicts_report(db: Session, user: User) -> Report:
    rows = db.scalars(select(Conflict).where(Conflict.workspace_id == user.workspace_id).order_by(Conflict.detected_at.desc())).all()
    mm = {m.id: m for m in db.scalars(select(Metric).where(Metric.workspace_id == user.workspace_id)).all()}
    cols = ["Metric key", "Reference", "Variant", "Type", "Severity", "Status", "Source", "Explanation", "Recommendation", "Resolution", "Detected"]
    data = [[
        c.metric_key, mm[c.metric_a_id].name + f" ({mm[c.metric_a_id].department or '-'})", mm[c.metric_b_id].name + f" ({mm[c.metric_b_id].department or '-'})",
        c.conflict_type, c.severity, c.status, "TARA suggestion" if c.conflict_type == "semantic" else "deterministic", c.explanation, c.recommendation,
        (c.resolution_action or "") + (f": {c.resolution_note}" if c.resolution_note else ""), fmt_datetime_ist(c.detected_at),
    ] for c in rows if c.metric_a_id in mm and c.metric_b_id in mm]
    return "Conflicts", cols, data


def evaluation_report(db: Session, user: User) -> Report:
    r = db.scalar(select(EvaluationRun).where(EvaluationRun.workspace_id == user.workspace_id).order_by(EvaluationRun.created_at.desc()).limit(1))
    cols = ["Measure", "Value", "Detail"]
    if not r:
        return "Evaluation results", cols, [["No evaluation has been run yet", "", ""]]
    x = r.results
    ca = x["conflict_detection"]
    data = [
        ["Run at", fmt_datetime_ist(r.created_at), f"{r.duration_ms:.0f} ms"],
        ["Dataset processing avg (ms)", x["dataset_processing"]["avg_ms"], f"{x['dataset_processing']['jobs']} jobs"],
        ["Query latency benchmark avg (ms)", x["query_latency"]["benchmark_avg_ms"], f"p95 {x['query_latency']['benchmark_p95_ms']}"],
        ["Validation pass rate", x["validation"]["pass_rate"], f"{x['validation']['metrics_passing']}/{x['validation']['metrics_checked']} metrics"],
        ["Query success rate", x["query_success"]["success_rate"], f"{x['query_success']['ok']}/{x['query_success']['total']}"],
        ["Conflict detection precision (automated)", ca["automated"]["precision"], f"{ca['cases']} labelled cases"],
        ["Conflict detection recall (automated)", ca["automated"]["recall"], ""],
        ["Conflict detection precision (manual text diff)", ca["naive_manual"]["precision"], ""],
        ["Conflict detection recall (manual text diff)", ca["naive_manual"]["recall"], ""],
        ["AI suggestion acceptance rate", x["ai_acceptance"]["acceptance_rate"], f"{x['ai_acceptance']['accepted']}/{x['ai_acceptance']['generations']}"],
        ["Dashboard load avg (ms)", x["dashboard_load"]["avg_ms"], f"{x['dashboard_load']['dashboards']} dashboards"],
    ]
    return "Evaluation results", cols, data


def dashboard_report(db: Session, user: User, d: Dashboard) -> Report:
    out = dashboard_data(db, d, user.id)
    cols = ["Widget", "Type", "Metric", "Dimension", "Value", "Unit", "Formatted value"]
    rows = []
    for w in d.widgets or []:
        res = out["widgets"].get(w["id"], {})
        if not res.get("ok"):
            rows.append([w.get("title"), w.get("type"), w.get("metric_key"), "", None, "", f"ERROR: {res.get('error')}"])
            continue
        unit = res.get("unit") or ""
        if w.get("type") == "kpi":
            v = res["rows"][0][-1] if res["rows"] else None
            rows.append([w.get("title"), "kpi", w.get("metric_key"), "", v, unit, format_value(v, unit)])
        else:
            for r in res["rows"]:
                rows.append([w.get("title"), w.get("type"), w.get("metric_key"), r[0], r[-1], unit, format_value(r[-1], unit)])
    return f"Dashboard: {d.name}", cols, rows


def to_csv(rep: Report) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(rep[1])
    for r in rep[2]:
        w.writerow(["" if v is None else v for v in r])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def to_xlsx(rep: Report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = rep[0][:31]
    ws.append(rep[1])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="111111")
    INR_FMT = r'[>=10000000]"₹"##\,##\,##\,##0.00;[>=100000]"₹"##\,##\,##0.00;"₹"##,##0.00'
    unit_idx = rep[1].index("Unit") if "Unit" in rep[1] else None
    val_idx = rep[1].index("Value") if "Value" in rep[1] else None
    for r in rep[2]:
        ws.append(["" if v is None else v for v in r])
        if unit_idx is not None and val_idx is not None and str(r[unit_idx]).lower() == "inr" and isinstance(r[val_idx], (int, float)):
            ws.cell(ws.max_row, val_idx + 1).number_format = INR_FMT
    for i, col in enumerate(rep[1], 1):
        width = max([len(str(col))] + [len(str(r[i - 1])) for r in rep[2][:200] if i - 1 < len(r)])
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(60, width + 2)
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(rep: Report) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm, title=rep[0])
    st = getSampleStyleSheet()
    small = st["BodyText"].clone("small")
    small.fontSize, small.leading = 7, 9
    story = [Paragraph(f"IntelliMetrics AI - {rep[0]}", st["Title"]), Paragraph(f"Generated {fmt_datetime_ist(datetime.utcnow())}", st["Normal"]), Spacer(1, 6 * mm)]
    data = [[Paragraph(f"<b>{c}</b>", small) for c in rep[1]]] + [[Paragraph(str("" if v is None else v)[:300].replace("&", "&amp;").replace("<", "&lt;"), small) for v in r] for r in rep[2][:500]]
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e5e5e5")), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#999999")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(t)
    doc.build(story)
    return buf.getvalue()
